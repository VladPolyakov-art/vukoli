"""Reading CVs and job postings from pasted text, files (PDF, DOCX, TXT, MD, HTML) and links."""
from __future__ import annotations

import base64
import html
import io
import ipaddress
import re
import socket
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx

MAX_FILE_BYTES = 8 * 1024 * 1024
MAX_PAGE_BYTES = 3 * 1024 * 1024
MAX_TEXT = 20000


class DocumentError(ValueError):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def strip_html(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style|noscript|svg|head)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr)>", "\n", raw)
    text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)


def _host_is_public(host: str) -> bool:
    """Refuse localhost and private, link-local or reserved addresses (SSRF guard)."""
    if not host or host.lower() in ("localhost",) or host.lower().endswith((".localhost", ".local", ".internal")):
        return False
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast \
                or ip.is_unspecified:
            return False
    return True


def check_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise DocumentError("url_invalid", "Use a link that starts with http:// or https://.")
    if not _host_is_public(parsed.hostname):
        raise DocumentError("url_not_allowed", "This link points to a local or private address.")
    return url


def read_url(url: str, *, transport: httpx.BaseTransport | None = None, resolve: bool = True) -> str:
    """Fetch a job posting page and return readable text. Follows up to 5 redirects, each checked."""
    current = url.strip()
    with httpx.Client(timeout=15, follow_redirects=False, transport=transport,
                      headers={"User-Agent": "InterviewCoach/0.1 (+local app)", "Accept-Language": "en"}) as client:
        for _ in range(6):
            if resolve:
                check_url(current)
            elif urlparse(current).scheme not in ("http", "https"):
                raise DocumentError("url_invalid", "Use a link that starts with http:// or https://.")
            try:
                with client.stream("GET", current) as response:
                    if response.is_redirect and response.headers.get("location"):
                        current = urljoin(current, response.headers["location"])
                        continue
                    response.raise_for_status()
                    body = b""
                    for chunk in response.iter_bytes():
                        body += chunk
                        if len(body) > MAX_PAGE_BYTES:
                            break
                    encoding = response.encoding or "utf-8"
            except httpx.HTTPError as exc:
                raise DocumentError("url_unreachable",
                                    "Could not open the link. Copy the job posting text and paste it instead.") from exc
            text = strip_html(body.decode(encoding, errors="replace"))
            if len(text) < 300:
                raise DocumentError("url_empty", "The page is almost empty (it may need a login). "
                                                 "Copy the job posting text and paste it instead.")
            return text[:MAX_TEXT]
    raise DocumentError("url_unreachable", "Too many redirects.")


def _docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_file(name: str, data_b64: str) -> str:
    try:
        data = base64.b64decode(data_b64, validate=True)
    except ValueError as exc:
        raise DocumentError("file_invalid", "The file could not be decoded.") from exc
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError("file_too_large", "The file is larger than 8 MB.")
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        try:
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages[:20])
        except Exception as exc:
            raise DocumentError("file_unreadable", "Could not read the PDF.") from exc
    elif suffix == ".docx":
        try:
            text = _docx_text(data)
        except Exception as exc:
            raise DocumentError("file_unreadable", "Could not read the Word document.") from exc
    elif suffix in (".html", ".htm"):
        text = strip_html(data.decode("utf-8", errors="replace"))
    elif suffix in (".txt", ".md", ".markdown", ""):
        text = data.decode("utf-8", errors="replace")
    else:
        raise DocumentError("file_unsupported", "Supported files: PDF, DOCX, TXT, MD and HTML.")
    text = text.strip()
    if len(text) < 100:
        raise DocumentError("file_empty", "The file has almost no text.")
    return text[:MAX_TEXT]


TERM = re.compile(r"\b(?:[A-Z][a-z]+[A-Z]\w*|[A-Z]{2,}[a-z]?s?|[A-Z][a-zA-Z0-9]*(?:\.js|\.io)|[A-Z][a-z]{2,})\b")
COMMON_CAPS = {"The", "This", "That", "With", "And", "For", "You", "Your", "Our", "We", "They", "What", "When",
               "Where", "Who", "How", "Why", "Experience", "Requirements", "Responsibilities", "About", "Role",
               "Team", "Work", "Skills", "Education", "Summary", "Senior", "Junior", "Lead", "Present", "January",
               "February", "March", "April", "May", "June", "July", "August", "September", "October",
               "November", "December", "Monday", "Friday", "Must", "Nice", "Have", "Plus", "Bonus"}


def known_terms(*texts: str, limit: int = 60) -> list[str]:
    """Product, tool and company names from the user's own CV / JD: never corrected as language errors."""
    seen: dict[str, None] = {}
    for text in texts:
        for match in TERM.findall(text or ""):
            if match not in COMMON_CAPS and match not in seen:
                seen[match] = None
            if len(seen) >= limit:
                return list(seen)
    return list(seen)
